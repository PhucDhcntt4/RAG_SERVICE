import csv
import io
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook


class OfficeExtractionError(ValueError):
    pass


def _append_checked(parts: list[str], value: str, max_chars: int) -> None:
    value = value.strip()

    if not value:
        return

    current_size = sum(len(part) + 2 for part in parts)

    if current_size + len(value) > max_chars:
        raise OfficeExtractionError(
            f"Nội dung sau khi trích xuất vượt quá {max_chars} ký tự"
        )

    parts.append(value)


def extract_docx(data: bytes, *, max_chars: int) -> str:
    document = Document(io.BytesIO(data))
    parts: list[str] = []

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()

            if not text:
                continue

            style_name = (block.style.name or "").lower()
            level = None

            if style_name.startswith("heading"):
                number = style_name.removeprefix("heading").strip()

                if number.isdigit():
                    level = min(max(int(number), 1), 6)

            if level is not None:
                text = f"{'#' * level} {text}"

            _append_checked(parts, text, max_chars)

        elif isinstance(block, Table):
            if not block.rows:
                continue

            headers = [
                cell.text.strip() or f"Cột {index + 1}"
                for index, cell in enumerate(block.rows[0].cells)
            ]

            table_lines = ["Bảng dữ liệu:"]

            for row_number, row in enumerate(block.rows[1:], start=2):
                values = [cell.text.strip() for cell in row.cells]

                fields = [
                    f"{header}: {value}"
                    for header, value in zip(headers, values)
                    if value
                ]

                if fields:
                    table_lines.append(
                        f"Dòng {row_number} | " + " | ".join(fields)
                    )

            _append_checked(parts, "\n".join(table_lines), max_chars)

    text = "\n\n".join(parts).strip()

    if not text:
        raise OfficeExtractionError("Không tìm thấy nội dung trong file DOCX")

    return text


def extract_xlsx(data: bytes, *, max_chars: int) -> str:
    workbook = load_workbook(
        io.BytesIO(data),
        read_only=True,
        data_only=True,
    )

    parts: list[str] = []

    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)

            try:
                first_row = next(rows)
            except StopIteration:
                continue

            headers = [
                str(value).strip() if value is not None else f"Cột {index + 1}"
                for index, value in enumerate(first_row)
            ]

            sheet_parts = [f"# Sheet: {worksheet.title}"]

            for row_number, row in enumerate(rows, start=2):
                fields = []

                for header, value in zip(headers, row):
                    if value is None or str(value).strip() == "":
                        continue

                    fields.append(f"{header}: {str(value).strip()}")

                if fields:
                    sheet_parts.append(
                        f"Dòng {row_number} | " + " | ".join(fields)
                    )

            if len(sheet_parts) > 1:
                _append_checked(
                    parts,
                    "\n\n".join(sheet_parts),
                    max_chars,
                )
    finally:
        workbook.close()

    text = "\n\n".join(parts).strip()

    if not text:
        raise OfficeExtractionError("Không tìm thấy dữ liệu trong file XLSX")

    return text


def extract_csv(data: bytes, *, max_chars: int) -> str:
    try:
        decoded = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise OfficeExtractionError(
            "CSV phải sử dụng mã hóa UTF-8"
        ) from exc

    sample = decoded[:8192]

    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(decoded), dialect)

    try:
        first_row = next(reader)
    except StopIteration as exc:
        raise OfficeExtractionError("File CSV không có dữ liệu") from exc

    headers = [
        value.strip() or f"Cột {index + 1}"
        for index, value in enumerate(first_row)
    ]

    file_parts = ["# Dữ liệu CSV"]

    for row_number, row in enumerate(reader, start=2):
        fields = []

        for header, value in zip(headers, row):
            value = value.strip()

            if value:
                fields.append(f"{header}: {value}")

        if fields:
            file_parts.append(
                f"Dòng {row_number} | " + " | ".join(fields)
            )

    text = "\n\n".join(file_parts).strip()

    if len(text) > max_chars:
        raise OfficeExtractionError(
            f"Nội dung CSV vượt quá {max_chars} ký tự"
        )

    if len(file_parts) == 1:
        raise OfficeExtractionError("File CSV không có dòng dữ liệu")

    return text


def extract_office_file(
    filename: str,
    data: bytes,
    *,
    max_chars: int,
) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix == ".docx":
        return extract_docx(data, max_chars=max_chars)

    if suffix == ".xlsx":
        return extract_xlsx(data, max_chars=max_chars)

    if suffix == ".csv":
        return extract_csv(data, max_chars=max_chars)

    raise OfficeExtractionError(f"Không hỗ trợ định dạng {suffix}")